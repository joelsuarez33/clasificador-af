"""Parsea Politica_AF.xlsx (hoja "Test Version"), genera data/catalogo.json y carga la tabla BigQuery `politica_af`.

Filas con código de clase válido (entero de 8 dígitos) -> catálogo.
Filas sin código -> encabezados de rubro: se descartan del catálogo, se loguean y se usan
para anotar la jerarquía (`rubro`) de cada clase.

Uso:
    python -m preprocess.load_catalogo            # JSON + BigQuery
    python -m preprocess.load_catalogo --sin-bq   # solo JSON (sin llamadas a GCP)
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import ConfigError, get_settings, verificar_python
from app.embeddings import normalizar_denominacion
from app.models import CatalogoItem

logger = logging.getLogger("preprocess.load_catalogo")

HOJA = "Test Version"
_CODIGO_RE = re.compile(r"^\d{8}$")
_PREFIJOS_NOTA = ("clave", "nota", "hinweis", "note")

COL_CODIGO, COL_DESCRIPCION, COL_EXPLICACION = 0, 1, 2


@dataclass(frozen=True)
class Encabezado:
    fila: int
    nivel: int  # 0 = sección, 1 = grupo (código de grupo tipo "0051000 000"), 2 = subrubro
    texto: str


@dataclass
class ResultadoParseo:
    items: list[CatalogoItem] = field(default_factory=list)
    encabezados: list[Encabezado] = field(default_factory=list)
    notas: list[tuple[int, str]] = field(default_factory=list)
    duplicados: list[int] = field(default_factory=list)


def parsear_codigo(valor: Any) -> int | None:
    """Devuelve el código si es un entero de exactamente 8 dígitos; None en otro caso."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor if 10_000_000 <= valor <= 99_999_999 else None
    if isinstance(valor, float):
        return parsear_codigo(int(valor)) if valor.is_integer() else None
    texto = str(valor).strip()
    return int(texto) if _CODIGO_RE.fullmatch(texto) else None


def _texto(valor: Any) -> str:
    return normalizar_denominacion(valor) if valor is not None else ""


def _sin_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)).lower()


def _celda(fila: Sequence[Any], idx: int) -> Any:
    return fila[idx] if idx < len(fila) else None


def _inicio_datos(filas: list[Sequence[Any]]) -> int:
    """Índice de la primera fila de datos: la siguiente a la cabecera 'Descripción', si existe."""
    for i, fila in enumerate(filas[:15]):
        if _sin_acentos(_texto(_celda(fila, COL_DESCRIPCION))) == "descripcion":
            return i + 1
    return 0


def parsear_filas(filas: Iterable[Sequence[Any]]) -> ResultadoParseo:
    filas = list(filas)
    resultado = ResultadoParseo()
    vistos: set[int] = set()
    seccion = grupo = subrubro = ""
    previa_fue_encabezado = False

    inicio = _inicio_datos(filas)
    for n, fila in enumerate(filas[inicio:], start=inicio + 1):
        crudo_codigo = _celda(fila, COL_CODIGO)
        descripcion = _texto(_celda(fila, COL_DESCRIPCION))
        codigo = parsear_codigo(crudo_codigo)

        if codigo is not None:
            previa_fue_encabezado = False
            if not descripcion:
                logger.warning("Fila %d: código %s sin descripción, se descarta", n, codigo)
                continue
            if codigo in vistos:
                logger.warning("Fila %d: código %s duplicado, se conserva la primera aparición", n, codigo)
                resultado.duplicados.append(codigo)
                continue
            vistos.add(codigo)
            resultado.items.append(
                CatalogoItem(
                    clase=codigo,
                    descripcion=descripcion,
                    explicacion=_texto(_celda(fila, COL_EXPLICACION)),
                    rubro=" > ".join(p for p in (grupo, subrubro) if p),
                )
            )
            continue

        if not descripcion:
            continue  # fila vacía o solo con columnas ignoradas

        if _sin_acentos(descripcion).startswith(_PREFIJOS_NOTA):
            resultado.notas.append((n, descripcion))
            logger.debug("Fila %d: nota ignorada: %s", n, descripcion[:80])
            continue

        if crudo_codigo not in (None, ""):
            # Código de grupo no válido como clase (p. ej. "0051000 000")
            nivel, grupo, subrubro = 1, descripcion, ""
        elif descripcion.isupper():
            nivel, seccion, grupo, subrubro = 0, descripcion, "", ""
        else:
            nivel = 2
            subrubro = f"{subrubro} > {descripcion}" if previa_fue_encabezado and subrubro else descripcion
        previa_fue_encabezado = nivel == 2
        resultado.encabezados.append(Encabezado(fila=n, nivel=nivel, texto=descripcion))
        logger.info("Fila %d: encabezado nivel %d: %s", n, nivel, descripcion)

    return resultado


def parsear_catalogo(path: Path, hoja: str = HOJA) -> ResultadoParseo:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de política: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise ValueError(f"La hoja {hoja!r} no existe en {path.name}. Hojas: {wb.sheetnames}")
        return parsear_filas(wb[hoja].iter_rows(values_only=True))
    finally:
        wb.close()


def guardar_json(items: list[CatalogoItem], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps([i.model_dump() for i in items], ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def cargar_a_bigquery(items: list[CatalogoItem], settings, clients) -> None:
    from google.cloud import bigquery

    bq = clients.bigquery
    dataset = bigquery.Dataset(settings.dataset_fq)
    dataset.location = settings.bq_location
    bq.create_dataset(dataset, exists_ok=True)

    ahora = datetime.now(timezone.utc).isoformat()
    filas = [{**i.model_dump(), "fecha_carga": ahora} for i in items]
    job = bq.load_table_from_json(
        filas,
        settings.tabla_politica_fq,
        job_config=bigquery.LoadJobConfig(
            schema=[
                bigquery.SchemaField("clase", "INT64", mode="REQUIRED"),
                bigquery.SchemaField("descripcion", "STRING", mode="REQUIRED"),
                bigquery.SchemaField("explicacion", "STRING"),
                bigquery.SchemaField("rubro", "STRING"),
                bigquery.SchemaField("fecha_carga", "TIMESTAMP"),
            ],
            write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        ),
    )
    job.result()
    logger.info("Cargadas %d clases en %s", len(filas), settings.tabla_politica_fq)


def main(argv: list[str] | None = None) -> int:
    verificar_python()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, help="Ruta a Politica_AF.xlsx (por defecto POLITICA_AF_XLSX)")
    parser.add_argument("--salida", type=Path, help="Ruta de catalogo.json (por defecto CATALOGO_JSON)")
    parser.add_argument("--sin-bq", action="store_true", help="No cargar a BigQuery")
    args = parser.parse_args(argv)

    try:
        settings = get_settings()
    except ConfigError as exc:
        logger.error("%s", exc)
        return 1

    xlsx = args.xlsx or settings.politica_af_xlsx
    salida = args.salida or settings.catalogo_path
    resultado = parsear_catalogo(xlsx)
    if not resultado.items:
        logger.error("No se encontraron clases válidas en %s", xlsx)
        return 1

    guardar_json(resultado.items, salida)
    logger.info(
        "Catálogo: %d clases, %d encabezados descartados, %d notas, %d duplicados -> %s",
        len(resultado.items), len(resultado.encabezados), len(resultado.notas), len(resultado.duplicados), salida,
    )

    if not args.sin_bq:
        from app.gcp import crear_clientes

        cargar_a_bigquery(resultado.items, settings, crear_clientes(settings))
    return 0


if __name__ == "__main__":
    sys.exit(main())
