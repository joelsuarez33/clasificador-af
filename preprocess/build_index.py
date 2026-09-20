"""Construye/actualiza la tabla BigQuery `af_historico_embeddings` y su VECTOR INDEX.

Lee AF_definitivos_creados.xlsx, limpia denominaciones, embebe con Vertex AI
(text-embedding-005, task_type=RETRIEVAL_DOCUMENT) y carga en BigQuery.

Incremental: cada fila tiene `row_id = sha256(clase|denominacion_limpia)`; solo se
embeben y cargan las filas cuyo row_id no existe en la tabla. Se carga por tramos,
así que una ejecución interrumpida se retoma donde quedó.

Uso:
    python -m preprocess.build_index                  # incremental
    python -m preprocess.build_index --prune          # además borra filas que ya no están en el Excel
    python -m preprocess.build_index --full-refresh   # recrea la tabla desde cero
    python -m preprocess.build_index --dry-run        # solo parsea y reporta (sin GCP)
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import sys
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import ConfigError, Settings, get_settings, verificar_python
from app.embeddings import MAX_BATCH, embed_textos, limpiar_denominacion, normalizar_denominacion

logger = logging.getLogger("preprocess.build_index")

COLUMNA_CLASE = "clase"
COLUMNA_DENOMINACION = "denominacion del activo fijo"
FILAS_MINIMAS_INDICE = 5000  # mínimo de BigQuery para CREATE VECTOR INDEX IVF


@dataclass(frozen=True)
class FilaHistorica:
    row_id: str
    clase: int
    denominacion: str
    denominacion_limpia: str


def _normalizar_encabezado(valor: Any) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    return " ".join("".join(c for c in texto if not unicodedata.combining(c)).lower().split())


def calcular_row_id(clase: int, denominacion_limpia: str) -> str:
    return hashlib.sha256(f"{clase}|{denominacion_limpia}".encode("utf-8")).hexdigest()


def preparar_filas(pares: Iterable[tuple[Any, Any]]) -> tuple[list[FilaHistorica], dict[str, int]]:
    """Limpia y deduplica pares (clase, denominación). Devuelve filas y estadísticas."""
    filas: dict[str, FilaHistorica] = {}
    stats = {"leidas": 0, "descartadas": 0, "duplicadas": 0}
    for clase_cruda, denominacion_cruda in pares:
        stats["leidas"] += 1
        try:
            clase = int(clase_cruda)
        except (TypeError, ValueError):
            stats["descartadas"] += 1
            continue
        denominacion = normalizar_denominacion(denominacion_cruda) if denominacion_cruda is not None else ""
        if not denominacion:
            stats["descartadas"] += 1
            continue
        limpia = limpiar_denominacion(denominacion)
        row_id = calcular_row_id(clase, limpia)
        if row_id in filas:
            stats["duplicadas"] += 1
            continue
        filas[row_id] = FilaHistorica(row_id, clase, denominacion, limpia)
    return list(filas.values()), stats


def leer_historico(path: Path) -> tuple[list[FilaHistorica], dict[str, int]]:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo histórico: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        hoja = wb.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        encabezado = [_normalizar_encabezado(c) for c in next(filas)]
        try:
            i_clase = encabezado.index(COLUMNA_CLASE)
            i_denom = encabezado.index(COLUMNA_DENOMINACION)
        except ValueError as exc:
            raise ValueError(
                f"{path.name}: se esperaban las columnas 'Clase' y 'Denominación del activo fijo', "
                f"encontradas {encabezado}"
            ) from exc
        pares = ((f[i_clase], f[i_denom]) for f in filas if f and len(f) > max(i_clase, i_denom))
        return preparar_filas(pares)
    finally:
        wb.close()


def _schema():
    from google.cloud import bigquery

    return [
        bigquery.SchemaField("row_id", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("clase", "INT64", mode="REQUIRED"),
        bigquery.SchemaField("denominacion", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("denominacion_limpia", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("embedding", "FLOAT64", mode="REPEATED"),
        bigquery.SchemaField("modelo_embedding", "STRING"),
        bigquery.SchemaField("fecha_carga", "TIMESTAMP"),
    ]


def asegurar_tabla(bq, settings: Settings, *, recrear: bool) -> None:
    from google.cloud import bigquery

    dataset = bigquery.Dataset(settings.dataset_fq)
    dataset.location = settings.bq_location
    bq.create_dataset(dataset, exists_ok=True)
    if recrear:
        logger.warning("--full-refresh: eliminando %s (y su índice vectorial)", settings.tabla_historico_fq)
        bq.delete_table(settings.tabla_historico_fq, not_found_ok=True)
    bq.create_table(bigquery.Table(settings.tabla_historico_fq, schema=_schema()), exists_ok=True)


def verificar_modelo(bq, settings: Settings) -> None:
    filas = bq.query(
        f"SELECT DISTINCT modelo_embedding AS m FROM `{settings.tabla_historico_fq}`"
    ).result()
    modelos = {f["m"] for f in filas} - {None}
    esperado = f"{settings.embedding_model}@{settings.embedding_dim}"
    if modelos and modelos != {esperado}:
        raise RuntimeError(
            f"La tabla tiene embeddings de {sorted(modelos)} y la configuración usa {esperado}. "
            "Ejecutá con --full-refresh para regenerar todo el índice."
        )


def ids_existentes(bq, settings: Settings) -> set[str]:
    filas = bq.query(f"SELECT row_id FROM `{settings.tabla_historico_fq}`").result()
    return {f["row_id"] for f in filas}


def podar(bq, settings: Settings, ids_vigentes: list[str]) -> int:
    from google.cloud import bigquery

    job = bq.query(
        f"DELETE FROM `{settings.tabla_historico_fq}` WHERE row_id NOT IN UNNEST(@ids)",
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ArrayQueryParameter("ids", "STRING", ids_vigentes)]
        ),
    )
    job.result()
    return job.num_dml_affected_rows or 0


def cargar_tramo(bq, settings: Settings, filas: list[FilaHistorica], vectores: dict[str, list[float]]) -> None:
    from google.cloud import bigquery

    ahora = datetime.now(timezone.utc).isoformat()
    modelo = f"{settings.embedding_model}@{settings.embedding_dim}"
    registros = [
        {
            "row_id": f.row_id,
            "clase": f.clase,
            "denominacion": f.denominacion,
            "denominacion_limpia": f.denominacion_limpia,
            "embedding": vectores[f.denominacion_limpia],
            "modelo_embedding": modelo,
            "fecha_carga": ahora,
        }
        for f in filas
    ]
    bq.load_table_from_json(
        registros,
        settings.tabla_historico_fq,
        job_config=bigquery.LoadJobConfig(
            schema=_schema(), write_disposition=bigquery.WriteDisposition.WRITE_APPEND
        ),
    ).result()


def crear_vector_index(bq, settings: Settings) -> None:
    total = next(iter(bq.query(f"SELECT COUNT(*) AS n FROM `{settings.tabla_historico_fq}`").result()))["n"]
    if total < FILAS_MINIMAS_INDICE:
        # BigQuery rechaza CREATE VECTOR INDEX (IVF) con < 5000 filas; VECTOR_SEARCH funciona igual por fuerza bruta.
        logger.warning(
            "La tabla tiene %d filas (< %d mínimo de BigQuery para un índice IVF): no se crea VECTOR INDEX; "
            "VECTOR_SEARCH hará búsqueda exacta.",
            total, FILAS_MINIMAS_INDICE,
        )
        return
    bq.query(
        f"""
        CREATE VECTOR INDEX IF NOT EXISTS `{settings.vector_index}`
        ON `{settings.tabla_historico_fq}`(embedding)
        OPTIONS(index_type = 'IVF', distance_type = 'COSINE')
        """
    ).result()
    logger.info("VECTOR INDEX %s asegurado sobre %s (%d filas)", settings.vector_index, settings.tabla_historico_fq, total)


def main(argv: list[str] | None = None) -> int:
    verificar_python()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, help="Ruta a AF_definitivos_creados.xlsx (por defecto AF_HISTORICO_XLSX)")
    parser.add_argument("--full-refresh", action="store_true", help="Recrea la tabla desde cero")
    parser.add_argument("--prune", action="store_true", help="Borra de BigQuery filas que ya no están en el Excel")
    parser.add_argument("--batch-size", type=int, default=100, help=f"Textos por request de embeddings (1-{MAX_BATCH})")
    parser.add_argument("--tramo", type=int, default=1000, help="Filas por carga a BigQuery")
    parser.add_argument("--dry-run", action="store_true", help="Solo parsea y reporta, sin llamar a GCP")
    args = parser.parse_args(argv)

    try:
        settings = get_settings()
    except ConfigError as exc:
        logger.error("%s", exc)
        return 1

    filas, stats = leer_historico(args.xlsx or settings.af_historico_xlsx)
    logger.info(
        "Histórico: %d leídas, %d válidas únicas, %d duplicadas, %d descartadas",
        stats["leidas"], len(filas), stats["duplicadas"], stats["descartadas"],
    )
    if args.dry_run:
        return 0
    if not filas:
        logger.error("No hay filas válidas para indexar")
        return 1

    from app.gcp import crear_clientes

    clients = crear_clientes(settings)
    bq = clients.bigquery
    asegurar_tabla(bq, settings, recrear=args.full_refresh)
    verificar_modelo(bq, settings)

    if args.prune and not args.full_refresh:
        borradas = podar(bq, settings, [f.row_id for f in filas])
        logger.info("Poda: %d filas eliminadas", borradas)

    existentes = set() if args.full_refresh else ids_existentes(bq, settings)
    nuevas = [f for f in filas if f.row_id not in existentes]
    logger.info("%d filas ya indexadas, %d nuevas a embeber", len(filas) - len(nuevas), len(nuevas))

    for inicio in range(0, len(nuevas), args.tramo):
        tramo = nuevas[inicio : inicio + args.tramo]
        textos = sorted({f.denominacion_limpia for f in tramo})
        vectores = embed_textos(
            clients.genai,
            textos,
            task_type="RETRIEVAL_DOCUMENT",
            model=settings.embedding_model,
            dim=settings.embedding_dim,
            batch_size=args.batch_size,
        )
        cargar_tramo(bq, settings, tramo, dict(zip(textos, vectores, strict=True)))
        logger.info("Cargadas %d/%d filas nuevas", inicio + len(tramo), len(nuevas))

    crear_vector_index(bq, settings)
    return 0


if __name__ == "__main__":
    sys.exit(main())
